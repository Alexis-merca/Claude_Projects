#!/usr/bin/env python3
"""Mercateam deck (base : build_swimlane_deck.py du skill mercateam-swimlane-deck)
+ types de slides : divider, stats, cards, timeline, chart, quotes.
usage: python3 build.py spec.json out.pptx --icons DIR --template template.pptx
"""
import json, re, argparse, os, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree

ap = argparse.ArgumentParser()
ap.add_argument('spec'); ap.add_argument('out')
ap.add_argument('--icons', default='icons'); ap.add_argument('--template')
A = ap.parse_args()
SPEC = json.load(open(A.spec, encoding='utf-8'))
IC = A.icons.rstrip('/') + '/'

rgb = RGBColor.from_string
PURPLE, TITLE, INK, GREY = rgb('5E4DA4'), rgb('6733FC'), rgb('1B1B3A'), rgb('6B6B80')
FONT, TFONT = SPEC.get('font', 'Inter'), SPEC.get('title_font', 'Inter')
STY = {'std': ('FFFFFF', '5E4DA4', None, 0.75, '5E4DA4'),
       'pain': ('FDEEEE', 'DC2626', None, 1.25, '7F1D1D'),
       'wip': ('FFF7E6', 'D97706', 'dash', 1.25, '5E4DA4'),
       'new': ('F1ECFF', '6733FC', None, 1.25, '5E4DA4'),
       'tool': ('F7F7FA', 'B9B3D6', 'dash', 0.75, '5E4DA4')}
LEG = {'pain': 'Point de douleur', 'std': 'Étape du process', 'wip': '! Pas encore en place, en cours',
       'new': 'Nouveau / proposé', 'tool': 'Outil / support'}
W, H = 13.333, 7.5

# ------------------------------------------------------------------ deck
prs = Presentation(A.template)
os.makedirs(IC, exist_ok=True)
for sl_ in prs.slides:   # pictos du template (slide ICONES)
    for sh in sl_.shapes:
        if sh.shape_type != 13:
            continue
        key = sh.name[5:] if sh.name.startswith('icon:') else None
        if not key:
            for t in sl_.shapes:
                if t.has_text_frame and 0 < t.top - sh.top < Inches(1.2) and abs(t.left - sh.left) < Inches(0.4) \
                        and re.fullmatch(r'[a-z_]+', t.text_frame.text.strip()):
                    key = t.text_frame.text.strip(); break
        if key:
            open(IC + key + '.png', 'wb').write(sh.image.blob)
lst = prs.slides._sldIdLst
for sid in list(lst)[1:]:
    prs.part.drop_rel(sid.rId); lst.remove(sid)
cov = SPEC.get('cover', {})
for sh in prs.slides[0].shapes:
    if not sh.has_text_frame:
        continue
    ps = sh.text_frame.paragraphs
    if any(re.search(r'\d{2}/\d{2}/\d{4}', p.text) for p in ps):
        for p in ps:
            if not p.runs:
                continue
            if re.search(r'\d{2}/\d{2}/\d{4}', p.text):
                p.runs[0].text = cov.get('date', p.runs[0].text)
                for r in p.runs[1:]: r.text = ''
            else:
                p.runs[0].text = cov.get('title', p.runs[0].text)
                for r in p.runs[1:]: r.text = ''
                br = p._p.find(qn('a:br'))
                if cov.get('subtitle') and br is not None:   # sous-titre en Inter Light après le saut de ligne
                    r2 = copy.deepcopy(p.runs[0]._r)
                    rpr = r2.find(qn('a:rPr')); rpr.set('b', '0')
                    for tag in ('a:latin', 'a:ea', 'a:cs', 'a:sym'):
                        e = rpr.find(qn(tag))
                        if e is not None: e.set('typeface', 'Inter Light')
                    r2.find(qn('a:t')).text = cov['subtitle']
                    br.addnext(r2)
names = [l.name for l in prs.slide_layouts]
layout = prs.slide_layouts[names.index('SECTION_HEADER')]


def dash(line):
    etree.SubElement(line._get_or_add_ln(), qn('a:prstDash')).set('val', 'dash')


def run_fmt(r, size, bold=False, color=INK, font=None, italic=False):
    f = r.font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = color; f.name = font or FONT


def text(s, x, y, w, h, t, size=9, bold=True, color=PURPLE, align=PP_ALIGN.CENTER, font=None,
         anchor=MSO_ANCHOR.MIDDLE, italic=False, spacing=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, line in enumerate(t if isinstance(t, list) else [t]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        if spacing: p.space_after = Pt(spacing)
        r = p.add_run(); r.text = line; run_fmt(r, size, bold, color, font, italic)
    return tb


def bullets(tf, items, size=10.5, color=INK, first=False, gap=4):
    for k, it in enumerate(items):
        p = tf.paragraphs[0] if (first and k == 0) else tf.add_paragraph()
        p.space_before = Pt(gap)
        pPr = p._p.get_or_add_pPr(); pPr.set('marL', str(Inches(0.16))); pPr.set('indent', str(-Inches(0.16)))
        bc = etree.SubElement(pPr, qn('a:buClr')); etree.SubElement(bc, qn('a:srgbClr')).set('val', '6733FC')
        etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
        etree.SubElement(pPr, qn('a:buChar')).set('char', '•')
        if isinstance(it, dict):   # {"b": "gras", "t": " suite"}
            r = p.add_run(); r.text = it['b']; run_fmt(r, size, True, color)
            r = p.add_run(); r.text = it.get('t', ''); run_fmt(r, size, False, color)
        else:
            r = p.add_run(); r.text = it; run_fmt(r, size, False, color)


def pic(s, name, x, y, size):
    path = IC + name + '.png'
    if os.path.exists(path):
        s.shapes.add_picture(path, Inches(x), Inches(y), Inches(size), Inches(size))


def rect(s, x, y, w, h, fill='FFFFFF', line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, dashed=False):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = rgb(fill)
    if line:
        b.line.color.rgb = rgb(line); b.line.width = Pt(lw)
        if dashed: dash(b.line)
    else:
        b.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        b.adjustments[0] = 0.06
    return b


def box(s, x, y, w, h, t, st):
    fill, line, d, lw, tc = STY[st]
    b = rect(s, x, y, w, h, fill, line, lw, dashed=bool(d))
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06); tf.margin_top = tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = t
    run_fmt(r, 9, True, rgb(tc))
    return b


def connect(s, a, b, dashed=False):
    same = abs((a.top + a.height / 2) - (b.top + b.height / 2)) < Emu(20000)
    if b.left >= a.left + a.width:
        i, j, k = 3, 1, (MSO_CONNECTOR.STRAIGHT if same else MSO_CONNECTOR.CURVE)
    elif a.left >= b.left + b.width:
        i, j, k = 1, 3, (MSO_CONNECTOR.STRAIGHT if same else MSO_CONNECTOR.CURVE)
    elif b.top > a.top:
        i, j, k = 2, 0, MSO_CONNECTOR.STRAIGHT
    else:
        i, j, k = 0, 2, MSO_CONNECTOR.STRAIGHT
    c = s.shapes.add_connector(k, 0, 0, 0, 0); c.begin_connect(a, i); c.end_connect(b, j)
    c.line.color.rgb = PURPLE; c.line.width = Pt(0.9)
    if dashed: dash(c.line)
    t = etree.SubElement(c.line._get_or_add_ln(), qn('a:tailEnd'))
    t.set('type', 'triangle'); t.set('w', 'med'); t.set('len', 'med')


def new_slide():
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    return s


def section(sl):
    s = new_slide()
    text(s, 0.34, 0.12, 10.6, 0.6, sl['title'], 24, color=TITLE, align=PP_ALIGN.LEFT, font=TFONT,
         anchor=MSO_ANCHOR.TOP)
    phase = sl.get('phase')
    if phase:
        hi = sl.get('phase_style', 'light' if phase.lower().startswith('avant') else 'accent')
        pill = rect(s, 10.9, 0.2, 2.1, 0.36, 'F3F3F3' if hi == 'light' else 'F1ECFF', shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        pill.adjustments[0] = 0.5
        tf = pill.text_frame; tf.margin_left = tf.margin_right = Inches(0.04)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = phase
        run_fmt(r, 10.5, True, rgb('555555') if hi == 'light' else TITLE)
    if sl.get('notes'):
        s.notes_slide.notes_text_frame.text = sl['notes']
    return s


def footer(s, sl, x=0.5, w=12.3, y=6.55):
    """message clé ou citation en bas de slide"""
    if sl.get('quote'):
        q = sl['quote']
        text(s, x, y - 0.12, 0.5, 0.6, '“', 40, color=TITLE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
        tb = s.shapes.add_textbox(Inches(x + 0.5), Inches(y), Inches(w - 0.5), Inches(0.75)); tf = tb.text_frame
        tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Inches(0.03)
        p = tf.paragraphs[0]; r = p.add_run(); r.text = q['text']; run_fmt(r, 11, False, INK, italic=True)
        r = p.add_run(); r.text = '  — ' + q['who']; run_fmt(r, 9.5, False, GREY)
    elif sl.get('message'):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.75)); tf = tb.text_frame
        tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP; tf.margin_left = Inches(0.03)
        p = tf.paragraphs[0]; r = p.add_run(); r.text = '→  '; run_fmt(r, 12, True, TITLE)
        r = p.add_run(); r.text = sl['message']; run_fmt(r, 12, True, INK)


# ------------------------------------------------------------------ slide types
def divider(sl):
    s = new_slide()
    text(s, 0.9, 2.0, 3, 1.4, sl['num'], 80, color=TITLE, align=PP_ALIGN.LEFT, font=TFONT, anchor=MSO_ANCHOR.BOTTOM)
    text(s, 0.9, 3.45, 10.5, 1.0, sl['title'], 34, color=INK, align=PP_ALIGN.LEFT, font=TFONT, anchor=MSO_ANCHOR.TOP)
    if sl.get('subtitle'):
        text(s, 0.9, 4.5, 9.5, 1.0, sl['subtitle'], 15, bold=False, color=PURPLE, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP)


def swimlane(sl):
    s = section(sl)
    lanes, panel = sl['lanes'], sl.get('panel')
    top, bot = 0.85, 6.3
    right = 9.55 if panel else 13.1
    ncol = max([n['col'] for n in sl['nodes']] + [0]) + 1
    x0, bw = 1.4, sl.get('box_w', 1.6)
    dx = min(1.95, (right - x0 - bw) / max(ncol - 1, 1)) if ncol > 1 else 1.95
    lh = (bot - top) / len(lanes); cy = []
    bh = min(0.9, lh - 0.25)
    for i, name in enumerate(lanes):
        c = top + lh * i + lh / 2; cy.append(c)
        pic(s, 'actor', 0.46, c - 0.36, 0.37)
        text(s, 0.02, c + 0.02, 1.25, 0.5, name, 8.5, anchor=MSO_ANCHOR.TOP)
        if i:
            l = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.3), Inches(top + lh * i), Inches(right), Inches(top + lh * i))
            l.line.color.rgb = rgb('ECE9F5'); l.line.width = Pt(0.5)
    v = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.25), Inches(top - .05), Inches(1.25), Inches(bot))
    v.line.color.rgb = rgb('D0CECE'); v.line.width = Pt(0.75); dash(v.line)
    shp, used = {}, set()
    for n in sl['nodes']:
        x = x0 + dx * n['col'] + n.get('dx', 0); y = cy[n['lane']] - bh / 2 + n.get('dy', 0)
        st = n.get('status', 'std'); used.add(st)
        shp[n['id']] = box(s, x, y, bw, bh, n['text'], st)
        if n.get('icon'): pic(s, n['icon'], x + bw - 0.16, y - 0.14, 0.26)
    for e in sl.get('edges', []):
        connect(s, shp[e[0]], shp[e[1]], len(e) > 2 and bool(e[2]))
    if panel:
        y = top
        for blk in panel:
            hgt = blk.get('height', 1.6)
            b = rect(s, 9.8, y, 3.3, hgt, blk.get('fill', 'FFFFFF'), '5E4DA4')
            tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = tf.margin_right = Inches(0.12); tf.margin_top = Inches(0.08)
            p = tf.paragraphs[0]; r = p.add_run(); r.text = blk['title']; run_fmt(r, 11, True, TITLE)
            bullets(tf, blk['items'], 9.5, PURPLE, gap=3)
            y += hgt + 0.2
    footer(s, sl, x=1.4, w=8.6, y=6.5)
    legend = sl.get('legend') or [k for k in ['pain', 'std', 'wip', 'new', 'tool'] if k in used]
    for k, st in enumerate(legend):
        fill, line, d, _, _ = STY[st]; yy = 6.55 + k * 0.21
        sw = rect(s, 10.4, yy + .03, .22, .13, fill, line, .9, dashed=bool(d))
        text(s, 10.68, yy - .02, 2.5, .22, sl.get('legend_labels', {}).get(st, LEG[st]), 8.5, bold=False,
             align=PP_ALIGN.LEFT)


def table_slide(sl):
    s = section(sl)
    rows = sl['rows']; cols = len(rows[0])
    widths = sl.get('widths') or [12.3 / cols] * cols
    tw = sum(widths)
    gt = s.shapes.add_table(len(rows), cols, Inches(0.5), Inches(0.95), Inches(tw),
                            Inches(0.4 * len(rows)))
    t = gt.table
    tblPr = gt._element.graphic.graphicData.tbl.tblPr
    for k in ('bandRow', 'firstRow'): tblPr.set(k, '0')
    st = tblPr.find(qn('a:tableStyleId'))
    if st is not None: st.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'   # No Style, Table Grid
    fs = sl.get('font_size', 10)
    for j in range(cols):
        t.columns[j].width = Inches(widths[j])
        for i in range(len(rows)):
            c = t.cell(i, j); val = rows[i][j]
            c.fill.solid(); c.fill.fore_color.rgb = rgb('6733FC' if i == 0 else ('FFFFFF' if i % 2 else 'F7F6FB'))
            c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.05)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; r = p.add_run(); r.text = str(val)
            hl = sl.get('highlight_col') == j and i > 0
            run_fmt(r, fs + (0.5 if i == 0 else 0), i == 0 or j == 0 or hl,
                    rgb('FFFFFF') if i == 0 else (TITLE if hl else INK))
            if hl: p.alignment = PP_ALIGN.CENTER
            for side in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):   # filets fins gris
                tcPr = c._tc.get_or_add_tcPr()
                ln = etree.SubElement(tcPr, qn(side)); ln.set('w', '6350')
                sf = etree.SubElement(ln, qn('a:solidFill')); etree.SubElement(sf, qn('a:srgbClr')).set('val', 'E4E0F2')
    if sl.get('row_heights'):
        for i, hh in enumerate(sl['row_heights']): t.rows[i].height = Inches(hh)
    footer(s, sl)


def stats_slide(sl):
    s = section(sl)
    st = sl['stats']; n = len(st)
    x0, y0, tw, gap = 0.5, 0.95, 12.3, 0.25
    w = (tw - gap * (n - 1)) / n; h = sl.get('stats_h', 1.35)
    for k, it in enumerate(st):
        x = x0 + k * (w + gap)
        rect(s, x, y0, w, h, 'F7F5FF', None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(s, x + 0.15, y0 + 0.08, w - 0.3, 0.62, it['value'], sl.get('value_size', 30), color=TITLE, font=TFONT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.BOTTOM)
        text(s, x + 0.15, y0 + 0.72, w - 0.3, h - 0.78, it['label'], 11.5, bold=False, color=INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    if sl.get('cards'):
        cards_block(s, sl, y0 + h + 0.35)
    if sl.get('chart'):
        chart_block(s, sl['chart'], 0.5, y0 + h + 0.3, sl['chart'].get('w', 12.3), 6.35 - (y0 + h + 0.3))
    footer(s, sl)


def cards_block(s, sl, y0, x0=0.5, tw=12.3):
    cards = sl['cards']; ncol = sl.get('cols', len(cards))
    nrow = (len(cards) + ncol - 1) // ncol
    gap = 0.25
    w = (tw - gap * (ncol - 1)) / ncol
    bottom = sl.get('cards_bottom', 6.3)
    h = (bottom - y0 - gap * (nrow - 1)) / nrow
    for k, c in enumerate(cards):
        r_, c_ = divmod(k, ncol)
        x = x0 + c_ * (w + gap); y = y0 + r_ * (h + gap)
        style = c.get('style', 'std')
        fill, line = {'std': ('FFFFFF', 'CFC8EC'), 'accent': ('F1ECFF', '6733FC'), 'pain': ('FFFFFF', 'E8B4B4')}[style]
        b = rect(s, x, y, w, h, fill, line, 1, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        stacked = w < 2.6
        tx = x + 0.2

        def badge(bx, by):
            if c.get('num'):
                o = rect(s, bx, by, 0.42, 0.42, '6733FC' if style != 'pain' else 'DC2626', shape=MSO_SHAPE.OVAL)
                tf = o.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = str(c['num'])
                run_fmt(r, 13, True, rgb('FFFFFF'), TFONT)
                return True
            if c.get('icon'):
                pic(s, c['icon'], bx, by, 0.42); return True
            return False
        tsize = sl.get('card_title_size', 14)
        if stacked:   # cartes étroites : pastille, puis titre pleine largeur
            badge(x + 0.2, y + 0.2)
            text(s, x + 0.15, y + 0.72, w - 0.3, 0.95, c['title'], tsize, color=INK, font=TFONT,
                 align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
            by = y + 1.7
        else:
            if badge(x + 0.2, y + 0.2): tx = x + 0.72
            text(s, tx, y + 0.14, x + w - tx - 0.12, 0.56, c['title'], tsize, color=INK,
                 font=TFONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
            by = y + 0.85
            if c.get('tag'):
                text(s, x + 0.2, y + 0.75, w - 0.4, 0.5, c['tag'], 10.5, color=TITLE, align=PP_ALIGN.LEFT,
                     anchor=MSO_ANCHOR.TOP)
                by = y + 1.3
        tb = s.shapes.add_textbox(Inches(x + 0.18), Inches(by), Inches(w - 0.36), Inches(y + h - by - 0.1))
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Inches(0.03); tf.margin_top = 0
        bullets(tf, c['items'], sl.get('card_size', 12.5), INK, first=True, gap=sl.get('card_gap', 8))


def cards_slide(sl):
    s = section(sl)
    if sl.get('intro'):
        text(s, 0.5, 0.82, 12.3, 0.5, sl['intro'], 12, bold=False, color=INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP)
    cards_block(s, sl, sl.get('cards_top', 1.45 if sl.get('intro') else 1.0))
    footer(s, sl)


def timeline_slide(sl):
    s = section(sl)
    ev = sl['events']; n = len(ev)
    x0, x1, ly = 1.35, 11.95, 3.72
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0 - 0.2), Inches(ly), Inches(x1 + 0.2), Inches(ly))
    ln.line.color.rgb = rgb('CFC8EC'); ln.line.width = Pt(2)
    step = (x1 - x0) / (n - 1); bw = min(2.35, step * 2 - 0.3)
    col = {'std': '6733FC', 'pain': 'DC2626', 'key': '1B1B3A'}
    for k, e in enumerate(ev):
        cx = x0 + k * step; st = e.get('status', 'std')
        d = rect(s, cx - 0.11, ly - 0.11, 0.22, 0.22, col[st], 'FFFFFF', 1.5, shape=MSO_SHAPE.OVAL)
        up = k % 2 == 0
        stem = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cx), Inches(ly - 0.12 if up else ly + 0.12),
                                      Inches(cx), Inches(ly - 0.45 if up else ly + 0.45))
        stem.line.color.rgb = rgb('CFC8EC'); stem.line.width = Pt(1)
        bx = min(max(cx - bw / 2, 0.4), W - 0.4 - bw)
        if up:
            text(s, bx, ly - 0.95, bw, 0.4, e['date'], 13.5, color=rgb(col[st]), font=TFONT, anchor=MSO_ANCHOR.BOTTOM)
            text(s, bx, ly - 2.75, bw, 1.8, e['text'], 11.5, bold=False, color=INK, anchor=MSO_ANCHOR.BOTTOM)
        else:
            text(s, bx, ly + 0.55, bw, 0.4, e['date'], 13.5, color=rgb(col[st]), font=TFONT, anchor=MSO_ANCHOR.TOP)
            text(s, bx, ly + 0.95, bw, 1.8, e['text'], 11.5, bold=False, color=INK, anchor=MSO_ANCHOR.TOP)
    if sl.get('phases'):   # bandeaux de phase sous la frise
        for ph in sl['phases']:
            xa = x0 + ph['from'] * step - 0.3; xb = x0 + ph['to'] * step + 0.3
            text(s, xa, 1.0 if False else 6.1, xb - xa, 0.3, ph['label'], 9, bold=False, color=GREY)
    footer(s, sl)


def chart_block(s, ch, x, y, w, h):
    cd = CategoryChartData(); cd.categories = ch['categories']
    for se in ch['series']:
        cd.add_series(se['name'], se['values'])
    kind = {'bar': XL_CHART_TYPE.BAR_CLUSTERED, 'col': XL_CHART_TYPE.COLUMN_CLUSTERED,
            'stack': XL_CHART_TYPE.COLUMN_STACKED}[ch.get('kind', 'col')]
    gf = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    c = gf.chart
    c.font.size = Pt(9.5); c.font.name = FONT; c.font.color.rgb = INK
    c.has_legend = len(ch['series']) > 1
    if c.has_legend:
        c.legend.position = XL_LEGEND_POSITION.TOP; c.legend.include_in_layout = False
        c.legend.font.size = Pt(10)
    if ch.get('title'):
        c.has_title = True; c.chart_title.text_frame.text = ch['title']
        r = c.chart_title.text_frame.paragraphs[0].runs[0]; run_fmt(r, 11, True, INK)
    else:
        c.has_title = False
    colors = ch.get('colors', ['6733FC', 'B9A6FF', 'DC2626'])
    plot = c.plots[0]; plot.gap_width = ch.get('gap', 60); plot.overlap = 100 if ch.get('kind') == 'stack' else -10
    for i, se in enumerate(plot.series):
        se.format.fill.solid(); se.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
    plot.has_data_labels = True
    dl = plot.data_labels; dl.font.size = Pt(9); dl.font.color.rgb = INK; dl.number_format = '0;;;'
    dl.number_format_is_linked = False
    if ch.get('kind') != 'stack':
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
    va = c.value_axis; va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb('ECE9F5'); va.format.line.fill.background()
    va.tick_labels.font.size = Pt(9); va.tick_labels.font.color.rgb = GREY
    ca = c.category_axis; ca.format.line.color.rgb = rgb('CFC8EC'); ca.tick_labels.font.size = Pt(9.5)
    if ch.get('reverse'):
        ca.reverse_order = True
    return gf


def chart_slide(sl):
    s = section(sl)
    ch = sl['chart']; cw = ch.get('w', 8.2)
    chart_block(s, ch, 0.4, 0.95, cw, 5.35)
    if sl.get('side'):
        x = 0.4 + cw + 0.3; w = 12.8 - x
        y = 1.05
        for blk in sl['side']:
            hh = blk.get('h', 1.0)
            b = rect(s, x, y, w, hh, 'F7F5FF', None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(s, x + 0.15, y + 0.06, w - 0.3, 0.55, blk['value'], 22, color=rgb(blk.get('color', '6733FC')),
                 font=TFONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
            text(s, x + 0.15, y + 0.58, w - 0.3, hh - 0.62, blk['label'], 11, bold=False, color=INK,
                 align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
            y += hh + 0.15
    footer(s, sl)


def quotes_slide(sl):
    s = section(sl)
    qs = sl['quotes']; ncol = sl.get('cols', 2)
    nrow = (len(qs) + ncol - 1) // ncol
    gap = 0.25; x0, y0, tw, bottom = 0.5, 1.0, 12.3, 6.3
    w = (tw - gap * (ncol - 1)) / ncol; h = (bottom - y0 - gap * (nrow - 1)) / nrow
    for k, q in enumerate(qs):
        r_, c_ = divmod(k, ncol); x = x0 + c_ * (w + gap); y = y0 + r_ * (h + gap)
        rect(s, x, y, w, h, 'F7F5FF', None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(s, x + 0.15, y + 0.02, 0.6, 0.7, '“', 44, color=TITLE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
        text(s, x + 0.7, y + 0.2, w - 0.9, h - 0.65, q['text'], sl.get('quote_size', 15), bold=False, color=INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, italic=True)
        text(s, x + 0.7, y + h - 0.45, w - 0.9, 0.35, q['who'], 10.5, bold=True, color=PURPLE, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP)
    footer(s, sl)


KIND = {'table': table_slide, 'divider': divider, 'stats': stats_slide, 'cards': cards_slide,
        'timeline': timeline_slide, 'chart': chart_slide, 'quotes': quotes_slide}
for sl in SPEC['slides']:
    KIND.get(sl.get('type'), swimlane)(sl)

for sl in list(prs.slides)[1:]:
    for el in list(sl.shapes._spTree.iter(qn('p:style'))):
        el.getparent().remove(el)
prs.save(A.out)
print('saved', A.out, len(prs.slides), 'slides')
